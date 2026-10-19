"""DOCX output and existing approval gate; no paid API calls."""
import copy
from io import BytesIO
import unittest
from unittest.mock import patch
from docx import Document
import exporting
import test_final_review as final_tests

class DocxExportTests(unittest.TestCase):
    for _name in ('setUp','tearDown','revisions','ai_response','run_mocked_ai','verified','plan','base','save_plan','body','draft_request','save_manual','final_payload','final_response','review','approve','approved','export'):
        locals()[_name] = getattr(final_tests.FinalReviewTests, _name)

    def test_downloads_and_approval_gate(self):
        self.assertEqual(self.export('docx').status_code, 409)
        self.approved()
        for fmt in ('html','json','docx'):
            self.assertEqual(self.export(fmt).status_code, 200)
        response = self.export('docx')
        self.assertIn('wordprocessingml.document', response.headers['content-type'])
        self.assertIn('.docx', response.headers['content-disposition'])
        document = Document(BytesIO(response.content))
        self.assertIn(self.article['draft']['content']['title'], [p.text for p in document.paragraphs])
        self.assertEqual(self.export('docx','wrong-approval').status_code,409)
        content = copy.deepcopy(self.article['draft']['content'])
        content['title'] += ' 수정'
        self.assertEqual(self.save_manual(content).status_code,200)
        self.assertEqual(self.export('docx').status_code,409)

    def test_content_order_guides_links_and_immutable_payload(self):
        self.approved()
        data = exporting.export_payload(self.article)
        data['article']['key_facts'] = [{'label':'선정 학교 수','value':'카드전용값','item_ids':[]}]
        data['article']['sections'] = [{'kind':k,'heading':'소제목'+str(i),'text':'본문'+str(i),'item_ids':[]} for i,k in enumerate(('fact','analysis','opinion','guidance','interview'))]
        data['article']['reader_guide'] = {'enabled':True,'audiences':[{'label':'학교','title':'가이드제목','sections':[{'type':'check','items':[{'text':'가이드내용'}]}]}]}
        before = copy.deepcopy(data)
        with patch('exporting.export_payload', return_value=data) as shared:
            blob = exporting.docx_export(self.article)
            shared.assert_called_once_with(self.article)
        self.assertEqual(data,before)
        doc = Document(BytesIO(blob)); text='\n'.join(p.text for p in doc.paragraphs)
        for label in ('사실·과거 사례·트렌드','편집자 분석','기자 의견·주장','학습 제안','인터뷰·취재'):
            self.assertIn(label,text)
        self.assertNotIn('카드전용값',text)
        positions=[text.index(x) for x in ('소제목0','본문4','가이드제목','가이드내용','원문과 연결한 근거','최종 검토 및 승인')]
        self.assertEqual(positions, sorted(positions))
        for enabled in (False, None):
            if enabled is None: data['article'].pop('reader_guide')
            else: data['article']['reader_guide']['enabled']=False
            with patch('exporting.export_payload',return_value=data):
                text='\n'.join(p.text for p in Document(BytesIO(exporting.docx_export(self.article))).paragraphs)
            self.assertNotIn('가이드내용',text)
