"""승인된 버전의 기사·출처·승인 메타데이터를 휴대 가능한 파일로 만든다."""
import html
import json
import re
from urllib.parse import urlsplit
from human_review import human_evidence_details
from evidence_catalog import evidence_catalog
from reader_guide import guide_items, LABELS


def export_payload(article):
    report = article["final_verification_report"]
    used = {citation["source_id"] for check in report["checks"] for citation in check["evidence"]}
    resolutions = article["approval"].get("resolutions", [])
    used.update(citation["source_id"] for item in resolutions for citation in item["evidence"])
    content = article["draft"]["content"]
    used_items = {item_id for block in content["sections"] + content["key_facts"] + content["audience_guidance"]
                  for item_id in block["item_ids"]}
    used_items.update(ref for item in guide_items(content.get("reader_guide")) for ref in item.item_ids)
    used.update(source_id for item in article["draft"]["metadata"].get("items", [])
                if item["id"] in used_items for source_id in item["source_ids"])
    snapshots = {source["source_id"]: source for source in report["read_sources"]}
    sources = [{"id": source["id"], "name": source["name"], "kind": source["kind"],
                "url": snapshots.get(source["id"], {}).get("resolved_url") or source["url"],
                "locator": source["locator"], "interview": source.get("interview"),
                "attachment_name": source["material"]["title"] if source.get("material") else None}
               for source in evidence_catalog(article) if source["id"] in used]
    return {"format_version": 1, "status": "approved_for_export", "article_id": article["id"],
            "article": article["draft"]["content"],
            "original": {"title": article["original_text"].splitlines()[0], "url": article["original_url"]},
            "sources": sources, "approval": article["approval"],
            "human_review": {"resolutions": resolutions, "evidence": human_evidence_details(report, resolutions),
                             "note": "AI 판정은 원본으로 보존하고, 항목별 사람의 근거 확인·안내 및 분석 판단을 승인 기록에 함께 저장했습니다."},
            "verification": {key: report[key] for key in
                ("id", "checked_at", "model", "draft_hash", "checks", "coverage_count", "reviewed_count", "scope")},
            "web_research": report.get("web_research", {"enabled": False, "status": "off"}),
            "attachment_note": "첨부 PDF 파일은 포함하지 않습니다. 게시할 때 파일을 별도로 첨부하거나 공개 출처 주소를 연결하세요."}


def json_export(article):
    return json.dumps(export_payload(article), ensure_ascii=False, indent=2)


def html_export(article):
    data = export_payload(article)
    content = data["article"]
    esc = lambda value: html.escape(str(value), quote=True)
    def link(text, url):
        if url and urlsplit(url).scheme in ("http", "https"):
            return '<a href="' + esc(url) + '" target="_blank" rel="noopener noreferrer">' + esc(text) + '</a>'
        return esc(text)
    audiences = {"student": "학생", "parent": "학부모", "teacher": "교사"}
    kinds = {"fact": "사실·과거 사례·트렌드", "analysis": "편집자 분석", "opinion": "기자 의견·주장", "guidance": "학습 제안", "interview": "인터뷰·취재"}
    body = ['<p class="category">' + esc(content["category"]) + '</p>', '<h1>' + esc(content["title"]) + '</h1>',
            '<div class="summary">' + ''.join('<p>' + esc(text) + '</p>' for text in content["summary"]) + '</div>']
    for section in content["sections"]:
        body.append('<section><span class="kind">' + esc(kinds[section["kind"]]) + '</span><h2>' + esc(section["heading"]) +
                    '</h2><p>' + esc(section["text"]) + '</p></section>')
    guide = content.get("reader_guide")
    if guide is None and content["audience_guidance"]:
        body.append('<h2>이 기사를 읽은 뒤 · 편집자의 제안</h2><div class="audiences">' + ''.join(
            '<section><strong>' + esc(audiences[item["audience"]]) + '</strong><h3>' + esc(item["heading"]) +
            '</h3><p>' + esc(item["text"]) + '</p></section>' for item in content["audience_guidance"]) + '</div>')
    if guide and guide["enabled"]:
        body.append('<section><h2>이 기사를 읽은 뒤 · 독자 가이드</h2>')
        for audience in guide["audiences"]:
            body.append('<section><h3>' + esc(audience["label"]) + '</h3><h4>' + esc(audience["title"]) + '</h4>')
            for section in audience["sections"]:
                body.append('<h4>' + esc(LABELS[section["type"]]) + '</h4><ul>' +
                            ''.join('<li>' + esc(item["text"]) + '</li>' for item in section["items"]) + '</ul>')
            body.append('</section>')
        body.append('</section>')
    original = data["original"]
    body.append('<section class="sources"><h2>원문과 연결한 근거</h2><p>원문: ' + link(original["title"], original["url"]) +
                ('' if original["url"] else ' · 주소 미등록') + '</p>')
    for source in data["sources"]:
        body.append('<p>' + link(source["name"], source["url"]) + (' · ' + esc(source["locator"]) if source["locator"] else '') +
                    (' · 첨부 자료: ' + esc(source["attachment_name"]) if source["attachment_name"] else '') + '</p>')
        if source.get("interview"):
            interview = source["interview"]
            body.append('<p>취재 기록: ' + esc(interview["speaker"]) + ' · ' + esc(interview.get("role", "")) +
                        ' · ' + esc(interview["interviewed_on"]) + ' · ' + esc(interview["method"]) + '</p>')
    body.append('</section>')
    approval = data["approval"]
    body.append('<footer>검토: ' + esc(approval["reviewer"]) + ' · 승인 시각: ' + esc(approval["approved_at"]) +
                ' · 기사 버전: ' + str(approval["draft_revision"]) +
                (' · 사람의 항목별 확인: ' + str(len(approval["resolutions"])) + '개' if approval.get("resolutions") else '') + '</footer>')
    return '<!doctype html><html lang="ko"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>' + esc(content["title"]) + '''</title>
<style>body{margin:0;background:#f5f7fb;color:#172b4a;font:16px/1.8 "Malgun Gothic",system-ui,sans-serif}article{max-width:900px;margin:30px auto;padding:32px;background:white;border-radius:12px}h1{font-size:30px;line-height:1.5}h2{font-size:20px;margin:28px 0 10px}h3{font-size:16px}p{white-space:pre-wrap;margin:8px 0}.summary{border-left:4px solid #245ddd;padding:12px 18px;background:#edf3ff}.audiences{display:grid;gap:12px;margin:24px 0}.audiences{grid-template-columns:repeat(3,1fr)}.audiences>section{background:#f5f7fb;border-radius:8px;padding:16px}.category,.kind,footer,.sources{font-size:13px;color:#61728b}a{color:#245ddd}.sources,footer{border-top:1px solid #dce4ef;padding-top:16px;margin-top:28px}@media(max-width:640px){article{margin:0;padding:20px;border-radius:0}.audiences{grid-template-columns:1fr}h1{font-size:26px}}</style></head><body><article>''' + ''.join(body) + '</article></body></html>'


XML_UNSAFE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def xml_safe(value):
    """Word cannot store control characters (e.g. vertical tabs) copied from PDFs."""
    if isinstance(value, str):
        return XML_UNSAFE.sub("", value)
    if isinstance(value, dict):
        return {key: xml_safe(item) for key, item in value.items()}
    if isinstance(value, list):
        return [xml_safe(item) for item in value]
    return value


def docx_export(article):
    """Build an editable document from the same approved payload as HTML/JSON."""
    from io import BytesIO
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.shared import Cm, Pt, RGBColor

    data = xml_safe(export_payload(article))
    content = data["article"]
    document = Document()
    page = document.sections[0]
    page.page_width, page.page_height = Cm(21), Cm(29.7)
    page.top_margin = page.bottom_margin = Cm(2)
    page.left_margin = page.right_margin = Cm(2.3)
    for name in ("Normal", "Title", "Heading 1", "Heading 2", "Heading 3", "List Bullet"):
        style = document.styles[name]
        style.font.name = "맑은 고딕"
        style.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "맑은 고딕")
        style.font.color.rgb = RGBColor(0, 0, 0)
        for border in list(style.element.xpath("./w:pPr/w:pBdr")):
            border.getparent().remove(border)
    document.styles["Normal"].font.size = Pt(11)
    document.styles["Normal"].paragraph_format.line_spacing = 1.3
    document.styles["Normal"].paragraph_format.space_after = Pt(7)
    document.styles["Title"].font.size = Pt(22)
    document.core_properties.title = content["title"]
    document.core_properties.author = data["approval"]["reviewer"]

    def add_link(paragraph, text, url):
        paragraph.add_run(text)
        if not url or urlsplit(url).scheme not in ("http", "https"):
            return
        paragraph.add_run(" · ")
        hyperlink = OxmlElement("w:hyperlink")
        hyperlink.set(qn("r:id"), paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True))
        run = OxmlElement("w:r")
        properties = OxmlElement("w:rPr")
        color = OxmlElement("w:color"); color.set(qn("w:val"), "245DDD")
        properties.append(color); run.append(properties)
        value = OxmlElement("w:t"); value.text = url
        run.append(value); hyperlink.append(run); paragraph._p.append(hyperlink)

    kinds = {"fact": "사실·과거 사례·트렌드", "analysis": "편집자 분석", "opinion": "기자 의견·주장",
             "guidance": "학습 제안", "interview": "인터뷰·취재"}
    document.add_paragraph(content["category"])
    document.add_paragraph(content["title"], "Title")
    for text in content["summary"]:
        document.add_paragraph(text)
    for section in content["sections"]:
        kind = document.add_paragraph(kinds[section["kind"]])
        kind.paragraph_format.keep_with_next = True
        document.add_heading(section["heading"], level=1)
        document.add_paragraph(section["text"])
    guide = content.get("reader_guide")
    if guide and guide["enabled"]:
        document.add_heading("이 기사를 읽은 뒤 · 독자 가이드", level=1)
        for audience in guide["audiences"]:
            document.add_heading(audience["label"], level=2)
            document.add_heading(audience["title"], level=3)
            for section in audience["sections"]:
                label = document.add_paragraph(LABELS[section["type"]])
                label.paragraph_format.keep_with_next = True
                for item in section["items"]:
                    document.add_paragraph(item["text"], "List Bullet")
    document.add_heading("원문과 연결한 근거", level=1)
    original = data["original"]
    add_link(document.add_paragraph(), "원문: " + original["title"] +
             (" · 주소 미등록" if not original["url"] else ""), original["url"])
    for source in data["sources"]:
        paragraph = document.add_paragraph()
        add_link(paragraph, source["name"], source["url"])
        if source["locator"]:
            paragraph.add_run(" · " + source["locator"])
        if source["attachment_name"]:
            paragraph.add_run(" · 첨부 자료: " + source["attachment_name"])
        if source.get("interview"):
            interview = source["interview"]
            document.add_paragraph("취재 기록: " + " · ".join([
                interview["speaker"], interview.get("role", ""), interview["interviewed_on"], interview["method"]]))
    approval = data["approval"]
    document.add_heading("최종 검토 및 승인", level=1)
    document.add_paragraph("검토: " + approval["reviewer"] + " · 승인 시각: " + approval["approved_at"] +
                           " · 기사 버전: " + str(approval["draft_revision"]))
    if approval.get("resolutions"):
        document.add_paragraph("사람의 항목별 확인: " + str(len(approval["resolutions"])) + "개")
    output = BytesIO()
    document.save(output)
    return output.getvalue()
